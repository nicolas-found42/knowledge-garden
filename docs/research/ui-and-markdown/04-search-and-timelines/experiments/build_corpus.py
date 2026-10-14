"""Build public/controlled research corpus; no private source paths or network access.
Use the bundled Python with reportlab and python-pptx. This is research data,
not a proposed durable product representation. Do not re-run over measured inputs.
"""
from pathlib import Path
import json, hashlib, shutil, wave, math, struct
from reportlab.pdfgen import canvas
from pptx import Presentation

R=Path(__file__).resolve().parent
A=R/'assets'; A.mkdir(parents=True,exist_ok=True)
ROOT=next(p for p in R.parents if (p/'AGENTS.md').exists())
OLD=ROOT/'docs/research/expanded-exploration/01-source-interpretation/experiments/media-lab'
for name in ['fruits.jpg','book-pdf-page-84.png','controlled-demo.mp4','controlled_demo.py','controlled-coding-thread.json']:
    shutil.copyfile(OLD/name,A/name)

controlled={
'picnic-caption.md':'# Picnic photo caption\nControlled metadata for the public fruit photograph, not its true provenance. Maya prepared fruit at Riverside on a day in May 2024. Photographer did not record the exact date. The photograph alone does not identify Maya or the location.\n',
'garden-download.html':'<!doctype html><meta charset="utf-8"><title>Garden visiting notes</title><h1>Garden visiting notes</h1><p>Controlled downloaded web page. Riverside glasshouse closes on Mondays. These notes describe a fictional visit, not a real venue schedule.</p>',
'field-notes.txt':'Controlled handwritten-note transcription surrogate. On 17 May 2024 at Riverside, Maya counted 12 pollinator visits in ten minutes. This was one observation, not a controlled trial. No handwritten original exists in this experiment.\n',
'visits.csv':'date,observer,site,count,minutes\n2024-05-17,Maya,Riverside,12,10\n2024-05-18,Maya,Riverside,8,10\n',
'receipt.txt':'Controlled ordinary download. Field notebook, one unit, 4.00 study tokens. Downloaded 2026-09-03. Transaction described: 2024-05-16. Not a real financial record.\n',
'plan-v1.md':'# Study plan, version 1\nControlled plan, 2024-05-18. Count visits at noon only. Maya proposed noon because the path is quieter. Proposal, not a measured best time.\n',
'plan-v2.md':'# Study plan, version 2\nControlled decision, 2024-05-20. Use morning and noon observations after rain disrupted the noon visit. This replaces the noon-only plan; better sampling accuracy has not been measured.\n',
'reasoning.md':'# Controlled study reasoning\n2024-05-18T09:00:00Z Maya proposes noon-only visits.\n2024-05-19T11:00:00Z Rain interrupts the noon visit.\n2024-05-20T08:00:00Z Maya adopts morning and noon observations. No accuracy comparison was run.\n',
'unknown-photo.txt':'# Missing source record\nControlled inaccessible original: winter-lake.heic was referenced but is not present. Caption mentions a blue canoe; date and photographer unknown. Do not treat a caption as inspected pixels.\n',
'conversation-link.md':'# Link occurrence\nControlled conversation mentioned https://example.invalid/field-guide. Acquisition: merely linked. Destination has never been fetched and its contents are unknown.\n',
'question.md':'# Open question\nDid rain change pollinator activity or only prevent observation? Current notes do not resolve this causal question.\n',
'trail.md':'# Study trail\nRead the visit notes, inspect counts, compare the two plans, then consider the unresolved rain question. This is an authored study route.\n',
}
for name,body in controlled.items(): (A/name).write_text(body)
pdfs={
'homework.pdf':['Controlled homework: sampling and uncertainty','Problem 1: 12 visits in 10 minutes gives 1.2 visits per minute.','One short observation does not estimate all-day activity.','Do not infer rain caused a change without comparable observations.'],
'field-report.pdf':['Controlled field report - Riverside','Maya recorded 12 pollinator visits in ten minutes on 17 May 2024.','The day was warm. A single observation does not establish a trend.','A later wet visit was interrupted, leaving missing observations.'],
}
for name,lines in pdfs.items():
    c=canvas.Canvas(str(A/name),invariant=1); c.setTitle(lines[0]); c.setFont('Helvetica',13)
    for i,line in enumerate(lines): c.drawString(40,800-i*30,line)
    c.save()
prs=Presentation()
for title,body in [('Field study briefing','Controlled presentation. Maya observed Riverside pollinators.'),('Observation, not causation','12 visits in ten minutes on 17 May 2024. Rain interrupted a later visit; no causal estimate.'),('Next sampling plan','Morning and noon. Compare against the earlier noon-only proposal.')]:
    s=prs.slides.add_slide(prs.slide_layouts[1]);s.shapes.title.text=title;s.placeholders[1].text=body
prs.save(A/'briefing.pptx')
with wave.open(str(A/'tone.wav'),'wb') as w:
    w.setparams((1,2,8000,0,'NONE','not compressed'))
    w.writeframes(b''.join(struct.pack('<h',int(4000*math.sin(2*math.pi*440*i/8000))) for i in range(8000)))

items=[]
def add(id,title,kind,fmt,body,file=None,origin='controlled',tags=(),activity='2026-09-29T20:00:00-04:00',start=None,end=None,label='Unknown',status='recorded',acq='captured'):
    href='/docs/research/ui-and-markdown/04-search-and-timelines/experiments/assets/'+file if file else ''
    items.append(dict(id=id,title=title,kind=kind,format=fmt,origin=origin,visibility='public' if origin.startswith('public') else 'controlled',body=body,tags=list(tags),sourceHref=href,activityTime=activity,eventStart=start,eventEnd=end,eventLabel=label,epistemic=status,acquisition=acq))
add('photo-fruit','Fruit on a picnic table','file','JPG','A real public OpenCV fruit photograph. Controlled picnic caption is separate; pixels do not establish Maya, Riverside, a date, nutrition, or weight.','fruits.jpg','public reused',('food','photo'))
add('caption','Picnic caption with an uncertain date','annotation','Markdown',controlled['picnic-caption.md'],'picnic-caption.md',tags=('Maya','Riverside','photo'),start='2024-05-01',end='2024-05-31',label='May 2024, day unknown')
add('homework','Sampling homework','file','PDF','Controlled homework: 12 visits in ten minutes is 1.2 per minute. A short observation does not estimate all-day activity. Rain causality is not established.','homework.pdf',tags=('learning','sampling','pollinators'))
add('report','Riverside field report','file','PDF','Maya recorded 12 pollinator visits in ten minutes on 17 May 2024 at Riverside. One observation does not establish a trend. A later wet visit was interrupted.','field-report.pdf',tags=('Maya','Riverside','pollinators'),start='2024-05-17',end='2024-05-17',label='17 May 2024')
add('slides','Field study briefing','file','PPTX','Controlled slides describe Maya at Riverside, 12 visits in ten minutes on 17 May 2024, interrupted wet visit and revised morning/noon plan.','briefing.pptx',tags=('learning','Maya','Riverside'),start='2024-05-21',end='2024-05-21',label='21 May 2024')
add('notes','Visit notes','file','TXT',controlled['field-notes.txt'],'field-notes.txt',tags=('Maya','Riverside','pollinators'),start='2024-05-17',end='2024-05-17',label='17 May 2024')
add('table','Two pollinator counts','file','CSV','Controlled rows: 17 May 2024, Maya, Riverside, 12 visits, 10 minutes; 18 May, 8 visits, 10 minutes. Rows do not attribute a difference to rain.','visits.csv',tags=('counts','pollinators'),start='2024-05-17',end='2024-05-18',label='17–18 May 2024')
add('download','Garden visiting notes','file','HTML','Controlled arbitrary download: fictional Riverside glasshouse closes on Mondays. This is not a real venue schedule.','garden-download.html',tags=('download','Riverside'))
add('receipt','Notebook purchase record','file','TXT',controlled['receipt.txt'],'receipt.txt',tags=('download','notebook'),activity='2026-09-03T09:00:00-04:00',start='2024-05-16',end='2024-05-16',label='Transaction 16 May 2024; imported later')
add('plan-1','Noon-only sampling proposal','decision','Markdown',controlled['plan-v1.md'],'plan-v1.md',tags=('sampling','Maya'),start='2024-05-18',end='2024-05-18',label='18 May 2024',status='tentative')
add('plan-2','Morning and noon sampling decision','decision','Markdown',controlled['plan-v2.md'],'plan-v2.md',tags=('sampling','Maya'),start='2024-05-20',end='2024-05-20',label='20 May 2024')
add('reasoning','Why the sampling plan changed','file','Markdown',controlled['reasoning.md'],'reasoning.md',tags=('sampling','rain'),start='2024-05-18',end='2024-05-20',label='18–20 May 2024')
add('person-maya','Maya','person','Markdown','Controlled named participant. Mentioned in caption, report, notes and presentation. No face identification has been run.',tags=('Maya',))
add('place-river','Riverside','place','Markdown','Fictional shared place in authored study records. It is not geolocated and is not a known real venue.',tags=('Riverside',))
add('concept-sampling','Sampling','concept','Markdown','Studying a limited set of observations. A common concept joins learning material, counts and plans.',tags=('learning','sampling'))
add('event-visit','The 17 May visit','event','Markdown','Controlled event described by notes, report and counts. Source import time is different from this described event.',tags=('Riverside',),start='2024-05-17',end='2024-05-17',label='17 May 2024')
add('claim-rate','Observed rate: 1.2 visits per minute','claim','Markdown','A calculation from 12 visits divided by ten minutes; bounded to the 17 May observation, not a daily estimate.',tags=('sampling','counts'))
add('claim-rain','Rain reduced pollinator activity','claim','Markdown','Speculative causal explanation, not established by the interrupted visit or the two rows. Distinguish missing observations from fewer insects.',tags=('rain','pollinators'),status='speculative')
add('question-rain','Did rain affect activity or observation?','question','Markdown',controlled['question.md'],'question.md',tags=('rain','gaps'))
add('trail','Study the observation before the explanation','trail','Markdown',controlled['trail.md'],'trail.md',tags=('learning',))
add('photo-missing','Blue canoe by a winter lake','file','HEIC',controlled['unknown-photo.txt'],'unknown-photo.txt',tags=('photo','missing'),acq='could_not_retrieve')
add('linked','Field-guide link occurrence','file','URL',controlled['conversation-link.md'],'conversation-link.md',tags=('link','gaps'),acq='merely_linked')
add('audio','One-second 440 Hz tone','file','WAV','Controlled waveform with no speech. This fixture checks retained file opening, not audio transcription or tone understanding.','tone.wav',tags=('audio',))
add('book','Historical book scan, page 84','file','PNG','Reused real public Darwin scan derivative. Printed-page and PDF-page numbering differ; see prior media manifest for original acquisition. No personal historical inspiration is implied.','book-pdf-page-84.png','public derivative',('book','learning'))
add('video','Motion difference demonstration','file','MP4','Reused executed demonstration built from public OpenCV video. At two seconds, 11566 is a sampled consecutive-frame changed-pixel count, not an interval aggregate or accuracy score.','controlled-demo.mp4','controlled over public video',('video','motion'))
add('code','Motion demonstration source','file','Python','Executed controlled code: grayscale, blur and threshold 35. It produces masks; better accuracy was not established.','controlled_demo.py',tags=('code','motion'))
add('chat','Controlled motion discussion','file','JSON','Authored conversation proposing and adopting a motion-mask change. This is not recovered personal history. Source code and measured output are separate records.','controlled-coding-thread.json',tags=('conversation','motion'))
relations=[]
def rel(fr,to,typ,ev,quote='',status='recorded',selector='document'):
    relations.append(dict(id='r'+str(len(relations)+1),**{'from':fr,'to':to},type=typ,status=status,evidence=[dict(itemId=ev,quote=quote,selector=selector)],selector=selector))
for s in ['caption','report','slides','notes']:
    rel(s,'person-maya','mentions',s,'Maya'); rel(s,'place-river','describes-place',s,'Riverside')
rel('caption','photo-fruit','annotates','caption','Controlled metadata','recorded')
for s in ['homework','report','table','plan-1','plan-2']: rel(s,'concept-sampling','about',s)
rel('notes','event-visit','describes','notes','17 May 2024')
rel('report','event-visit','describes','report','17 May 2024')
rel('table','claim-rate','supports','table','12 visits, 10 minutes',selector='row 1')
rel('homework','claim-rate','explains','homework','1.2 per minute',selector='page 1')
rel('plan-2','plan-1','revises','plan-2','This replaces the noon-only plan')
rel('reasoning','plan-2','explains','reasoning','Rain interrupts the noon visit')
rel('claim-rain','question-rain','needs-evidence','question-rain',status='speculative')
rel('table','claim-rain','insufficient-for','table','Rows do not attribute a difference to rain')
for s in ['notes','table','plan-1','plan-2','question-rain']: rel('trail',s,'trail-step','trail')
rel('code','video','produces','code',selector='source code')
rel('chat','code','motivates','chat')
rel('book','concept-sampling','possible-analogy','book',status='speculative')
reasoning=[dict(id='reason-'+str(i),time=t,type=typ,itemId=it,label=label,evidence=[dict(itemId='reasoning',selector='line '+str(i+1))]) for i,(t,typ,it,label) in enumerate([
('2024-05-18T09:00:00Z','proposal','plan-1','Propose noon-only visits'),('2024-05-19T11:00:00Z','observation','reasoning','Rain interrupts a visit'),('2024-05-20T08:00:00Z','revision','plan-2','Adopt morning and noon')],1)]
tasks=[
('explore','Follow a useful connection','Explore how a photograph annotation leads to a person, report and learning material.',['caption','person-maya','report','homework']),
('explain','Check the rate claim','Recover the calculation and its single-observation qualification.',['claim-rate','table','homework']),
('orient','Orient across the whole collection','Show files and extracted items, then expand a group without losing collection access.',['photo-fruit','person-maya','video']),
('remember','Recover a partly remembered file','Find the slides mentioning Maya at Riverside.',['slides']),
('compare','Compare plans','Explain noon-only versus morning-and-noon; identify what was never measured.',['plan-1','plan-2','reasoning']),
('change','Separate three kinds of time','Inspect activity, reasoning and described events; retain unknown and month-only dates.',['receipt','caption','reasoning']),
('gaps','Inspect gaps','Find the missing photograph, unacquired link and unsupported rain explanation.',['photo-missing','linked','question-rain','claim-rain']),
('trails','Follow the study trail','Follow evidence, versions and unresolved question in order.',['trail','notes','table','plan-1','plan-2','question-rain'])]
data=dict(schema='kg-research-corpus-v1',researchDate='2026-09-29',boundary='Public and explicitly controlled research only. Relationships are researcher-authored, not inferred or user accepted.',items=items,relations=relations,reasoning=reasoning,tasks=[dict(id=i,title=t,instruction=d,expectedIds=e) for i,t,d,e in tasks])
(R/'corpus.json').write_text(json.dumps(data,indent=2)+'\n')
manifest=[]
for p in sorted(A.iterdir()):manifest.append(dict(file='assets/'+p.name,bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest(),origin='reused see prior media-lab/sources.json' if p.name in ['fruits.jpg','book-pdf-page-84.png','controlled-demo.mp4','controlled_demo.py','controlled-coding-thread.json'] else 'controlled authored fixture'))
(R/'source-manifest.json').write_text(json.dumps(dict(files=manifest,previousManifest='../../../expanded-exploration/01-source-interpretation/experiments/media-lab/sources.json',conversationDestinationFetches=0),indent=2)+'\n')
print(json.dumps(dict(items=len(items),relations=len(relations),assets=len(manifest),tasks=len(tasks))))
